
from io import BytesIO

from docx import Document
from docx.shared import Inches
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
)
from reportlab.lib.units import inch
from xml.sax.saxutils import escape


def format_docx(
    content: str,
    document_type: str,
    terms=None,
    logo_bytes=None,
) -> bytes:
    doc = Document()

    if logo_bytes:
        image_stream = BytesIO(logo_bytes)
        doc.add_picture(image_stream, width=Inches(1.5))

    doc.add_heading(document_type, level=1)

    for paragraph in content.split("\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def format_pdf(
    content: str,
    document_type: str,
    logo_bytes=None,
) -> bytes:
    output = BytesIO()

    pdf = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=inch,
        bottomMargin=inch,
    )

    styles = getSampleStyleSheet()
    story = []

    story.append(
        Paragraph(escape(document_type), styles["Title"])
    )
    story.append(Spacer(1, 12))

    for paragraph in content.split("\n"):
        if paragraph.strip():
            safe_text = escape(paragraph)
            story.append(
                Paragraph(safe_text, styles["BodyText"])
            )
            story.append(Spacer(1, 8))

    pdf.build(story)